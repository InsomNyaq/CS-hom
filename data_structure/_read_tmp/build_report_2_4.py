# -*- coding: utf-8 -*-
"""基于 报告模板25级.docx 生成实验 2.4 的实验报告。"""
import shutil
import copy

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

TEMPLATE = r"D:\CS-hom\data_structure\lab\报告模板25级.docx"
TARGET = r"D:\CS-hom\data_structure\lab\lab02\实验报告_2.4_两个有序链表序列的合并.docx"
IMAGE = r"D:\CS-hom\data_structure\lab\lab02\merge.png"

BODY_FONT = "宋体"
BODY_COLOR = RGBColor(0x40, 0x40, 0x40)

CODE = r"""#include <stdio.h>
#include <stdlib.h>
typedef struct Node{
    int data;
    struct Node * next;
}Node;

//create a link
Node * createLink(){
    Node * head = malloc(sizeof(Node));
    head->next = NULL;
    Node * tail = head;
    int x;

    while(scanf("%d", &x) == 1 && x != -1){
        Node * p = malloc(sizeof(Node));
        p->data = x;
        p->next = NULL;

        tail->next = p;
        tail = p;
    }

    return head;
}

//merge two links
Node * mergeLink(Node* L1, Node* L2){
    Node * p = L1->next;
    Node * q = L2->next;

    Node * head = malloc(sizeof(Node));
    head->next = NULL;
    Node * tail = head;
    
    while(p != NULL && q !=NULL){
        if(p->data <= q->data){
            tail->next = p;
            p = p->next;
        }else{
            tail->next = q;
            q = q->next;
        }

        tail = tail->next;
    }

    if (p != NULL){
        tail->next = p;
    }else{
        tail->next = q;
    }

    return head;
}

//print the link_merged
void printLink(Node * L){
    Node * p = L->next;
    if(p == NULL){
        printf("NULL\n");
        return;
    }

    int first = 1;
    while(p)
    {
     if (!first){
        printf(" ");
    }

    printf("%d", p->data);
    p = p->next;
    first = 0;
    }
}


int main()
{
    Node * S1 = createLink();
    Node * S2 = createLink();

    Node * Lmerged = mergeLink(S1, S2);
    printLink(Lmerged);

    return 0;
}"""


def set_east_asia(run, name):
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def no_snap(p):
    """关闭“对齐到网格”，避免 12pt/代码行被按文档网格放大行距。"""
    pPr = p._p.get_or_add_pPr()
    el = pPr.find(qn("w:snapToGrid"))
    if el is None:
        el = pPr.makeelement(qn("w:snapToGrid"), {})
        pPr.append(el)
    el.set(qn("w:val"), "0")


def add_p(doc, text="", style=None, size=None, bold=None, font=None,
          color=None, align=None, first_indent=None, compact=False):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        if size is not None:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color
        if font is not None:
            run.font.name = font
            set_east_asia(run, font)
    if align is not None:
        p.alignment = align
    if first_indent is not None:
        p.paragraph_format.first_line_indent = Pt(first_indent)
    if compact:
        no_snap(p)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
    return p


def add_code(doc, text):
    for line in text.split("\n"):
        p = doc.add_paragraph(style="代码样式")
        no_snap(p)
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        if line:
            run = p.add_run(line)
            run.font.size = Pt(10)
            run.font.name = "Consolas"


def main():
    shutil.copyfile(TEMPLATE, TARGET)
    doc = Document(TARGET)
    body = doc.element.body

    # 清空正文，保留节属性（页面大小、页边距）
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)

    # 标题
    add_p(doc, "数 据 结 构 实 验 报 告", style="题目样式", size=22, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(doc, "", style="题目样式")
    add_p(doc, "实验 2.4 两个有序链表序列的合并（15分 2学时）", style="题目样式",
          size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # 个人信息
    add_p(doc, "姓名：                                   电子邮箱地址：",
          style="题目样式", size=10.5, first_indent=10.5)
    add_p(doc, "年级专业：26级计算机班", style="题目样式", size=10.5, first_indent=10.5)
    add_p(doc, "", style="题目样式")
    add_p(doc, "实验内容", style="题目样式", size=14, bold=True)

    # 题目
    add_p(doc, "两个有序链表序列的合并（15分 2学时）", size=10.5, bold=True,
          align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p(doc, "已知两个非降序链表序列S1与S2，设计函数构造出S1与S2合并后的新的非降序链表S3。",
          size=10.5, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=21)
    add_p(doc, "输入格式：", size=10.5, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p(doc, "输入分两行，分别在每行给出由若干个正整数构成的非降序序列，用−1表示序列的结尾"
              "（−1不属于这个序列）。数字用空格间隔。",
          size=10.5, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=21)
    add_p(doc, "输出格式：", size=10.5, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p(doc, "在一行中输出合并后新的非降序链表，数字间用空格分开，结尾不能有多余空格；"
              "若新链表为空，输出NULL。",
          size=10.5, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=21)
    add_p(doc, "输入样例：", size=10.5, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p(doc, "1 3 5 -1", size=12, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=21, compact=True)
    add_p(doc, "2 4 6 8 10 -1", size=12, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=21, compact=True)
    add_p(doc, "输出样例：", size=10.5, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p(doc, "1 2 3 4 5 6 8 10", size=12, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=21, compact=True)

    # 算法思路
    add_p(doc, "算法思路：", size=10.5, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p(doc, "两个链表都带表头结点，p、q 分别指向 S1、S2 的第一个数据结点。比较 p->data 与 "
              "q->data，把较小者尾插到新链表 S3 的末尾，并后移相应的指针；当其中一个链表先"
              "遍历完时，把另一个链表剩余的整段直接链接到 S3 的末尾。整个过程只修改指针，"
              "不申请新的数据结点，时间复杂度为 O(m+n)，空间复杂度为 O(1)。",
          size=10.5, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=21)

    # 代码
    p = add_p(doc, "代码", style="题目说明样式", size=10.5)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.page_break_before = True
    add_code(doc, CODE)

    # 运行结果
    add_p(doc, "运行结果：", style="题目说明样式", size=10.5).paragraph_format.space_before = Pt(10)
    add_p(doc, "", size=10.5)
    add_p(doc, "输入数据：", size=10.5)
    add_p(doc, "1 3 5 -1", size=12, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=21, compact=True)
    add_p(doc, "2 4 6 8 10 -1", size=12, font=BODY_FONT, color=BODY_COLOR,
          align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=21, compact=True)
    add_p(doc, "输出结果", size=10.5)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.add_run().add_picture(IMAGE, width=Inches(5.5))

    doc.save(TARGET)
    print("saved", TARGET)
    print("paragraphs", len(doc.paragraphs))


if __name__ == "__main__":
    main()

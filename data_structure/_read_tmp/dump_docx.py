# -*- coding: utf-8 -*-
import sys, io
from docx import Document
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.document import Document as _Document
from docx.section import Section


def iter_block_items(parent):
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        parent_elm = parent._element
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def para_text(p):
    return "".join(r.text for r in p.runs) if p.runs else p.text


def dump_table(t, indent, out):
    out.write("%s[TABLE %d rows x %d cols]\n" % (indent, len(t.rows), len(t.columns)))
    for ri, row in enumerate(t.rows):
        cells = []
        seen = set()
        for ci, cell in enumerate(row.cells):
            if id(cell._tc) in seen:
                cells.append("<merged>")
                continue
            seen.add(id(cell._tc))
            txt = " / ".join(para_text(p) for p in cell.paragraphs if para_text(p).strip())
            for sub in cell.tables:
                txt += " <nested table>"
            cells.append(txt.replace("\n", " "))
        out.write("%s  R%d: %s\n" % (indent, ri + 1, " | ".join(cells)))


def dump(path, outpath):
    doc = Document(path)
    with io.open(outpath, "w", encoding="utf-8") as out:
        out.write("FILE: %s\n" % path)
        cp = doc.core_properties
        out.write("title=%r subject=%r author=%r\n" % (cp.title, cp.subject, cp.author))
        out.write("=== SECTIONS ===\n")
        for si, sec in enumerate(doc.sections):
            out.write("section %d: page %s x %s, margins L%.0f R%.0f T%.0f B%.0f\n" % (
                si, sec.page_width, sec.page_height,
                sec.left_margin or 0, sec.right_margin or 0,
                sec.top_margin or 0, sec.bottom_margin or 0))
            for hname, hf in (("header", sec.header), ("footer", sec.footer),
                              ("first_header", sec.first_page_header), ("first_footer", sec.first_page_footer)):
                txts = [para_text(p) for p in hf.paragraphs if para_text(p).strip()]
                if txts:
                    out.write("  %s: %s\n" % (hname, " | ".join(txts)))
                for t in hf.tables:
                    out.write("  %s table:\n" % hname)
                    dump_table(t, "    ", out)
        out.write("=== BODY ===\n")
        i = 0
        for block in iter_block_items(doc):
            if isinstance(block, Paragraph):
                txt = para_text(block)
                style = block.style.name if block.style is not None else ""
                has_img = "graphicData" in block._p.xml
                if not txt.strip() and not has_img:
                    continue
                tag = []
                if style and style != "Normal":
                    tag.append("style=%s" % style)
                if has_img:
                    tag.append("IMG")
                if block.alignment is not None:
                    tag.append("align=%s" % block.alignment)
                i += 1
                out.write("P%03d [%s] %s\n" % (i, ", ".join(tag) if tag else "-", txt))
            else:
                i += 1
                dump_table(block, "", out)
        # styles summary
        out.write("=== STYLES USED ===\n")
        used = {}
        for p in doc.paragraphs:
            n = p.style.name if p.style is not None else "?"
            used[n] = used.get(n, 0) + 1
        for k, v in sorted(used.items(), key=lambda x: -x[1]):
            out.write("%s: %d\n" % (k, v))


if __name__ == "__main__":
    dump(sys.argv[1], sys.argv[2])

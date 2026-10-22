# -*- coding: utf-8 -*-
"""Dump paragraphs, tables and cell shading of the lab report."""
import sys
from docx import Document
from docx.document import Document as DocumentClass
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P

SRC = sys.argv[1]
OUT = sys.argv[2]

doc = Document(SRC)


def iter_block_items(parent):
    if isinstance(parent, DocumentClass):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def para_text(p):
    return p.text.replace("\n", "\\n")


def cell_shading(cell):
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return ""
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        return ""
    return "%s/%s" % (shd.get(qn("w:val")), shd.get(qn("w:fill")))


def cell_borders(cell):
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return ""
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        return ""
    parts = []
    for edge in ("top", "left", "bottom", "right"):
        el = borders.find(qn("w:" + edge))
        if el is not None:
            parts.append("%s=%s/%s" % (edge, el.get(qn("w:val")), el.get(qn("w:sz"))))
    return ",".join(parts)


lines = []
lines.append("=== SECTION ===")
for s in doc.sections:
    lines.append("page %sx%s margins L%s R%s T%s B%s" % (
        s.page_width, s.page_height, s.left_margin, s.right_margin,
        s.top_margin, s.bottom_margin))

lines.append("")
lines.append("=== BODY BLOCKS ===")
pi = 0
ti = 0
for block in iter_block_items(doc):
    if isinstance(block, Paragraph):
        p = block
        style = p.style.name if p.style is not None else ""
        marked = "  <<<" if ("{{" in p.text or "＿＿" in p.text or "____" in p.text) else ""
        lines.append("[P%03d][%s] %s%s" % (pi, style, para_text(p), marked))
        pi += 1
    else:
        t = block
        lines.append("[T%02d] style=%s rows=%d cols=%d" % (
            ti, t.style.name if t.style is not None else "", len(t.rows), len(t.columns)))
        for ri, row in enumerate(t.rows):
            seen = set()
            for ci, cell in enumerate(row.cells):
                key = id(cell._tc)
                if key in seen:
                    continue
                seen.add(key)
                txt = cell.text.replace("\n", " | ")
                lines.append("    (r%d,c%d) shd=[%s] borders=[%s] text=%r" % (
                    ri, ci, cell_shading(cell), cell_borders(cell), txt))
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.text.strip() or r.font.color and r.font.color.rgb:
                            pass
        ti += 1

lines.append("")
lines.append("=== HEADERS/FOOTERS ===")
for si, s in enumerate(doc.sections):
    for name in ("header", "footer", "first_page_header", "first_page_footer"):
        try:
            part = getattr(s, name)
        except Exception:
            continue
        if part is None:
            continue
        txt = " | ".join(p.text for p in part.paragraphs if p.text.strip())
        if txt:
            lines.append("sec%d %s: %s" % (si, name, txt))

lines.append("")
lines.append("=== STYLES USED ===")
used = sorted({p.style.name for p in doc.paragraphs if p.style is not None})
lines.append(", ".join(used))

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("wrote", OUT, len(lines), "lines")


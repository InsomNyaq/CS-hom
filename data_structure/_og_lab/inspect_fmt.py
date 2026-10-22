# -*- coding: utf-8 -*-
"""Inspect paragraph shading/borders, numbering and images."""
import sys
from docx import Document
from docx.oxml.ns import qn

SRC = sys.argv[1]
OUT = sys.argv[2]
doc = Document(SRC)
out = []

out.append("=== IMAGE RELS ===")
for rel in doc.part.rels.values():
    if "image" in rel.reltype:
        out.append("%s %s %s" % (rel.rId, rel.target_ref, rel.reltype))

out.append("")
out.append("=== PARAGRAPH DETAIL ===")
for i, p in enumerate(doc.paragraphs):
    pPr = p._p.pPr
    shd = pPr.find(qn("w:shd")) if pPr is not None else None
    pbdr = pPr.find(qn("w:pBdr")) if pPr is not None else None
    numPr = pPr.find(qn("w:numPr")) if pPr is not None else None
    info = []
    if shd is not None:
        info.append("shd=%s/%s" % (shd.get(qn("w:val")), shd.get(qn("w:fill"))))
    if pbdr is not None:
        edges = []
        for edge in ("top", "left", "bottom", "right"):
            el = pbdr.find(qn("w:" + edge))
            if el is not None:
                edges.append("%s:%s/%s/%s" % (edge, el.get(qn("w:val")), el.get(qn("w:sz")), el.get(qn("w:color"))))
        info.append("bdr[" + ",".join(edges) + "]")
    if numPr is not None:
        ilvl = numPr.find(qn("w:ilvl"))
        numId = numPr.find(qn("w:numId"))
        info.append("num=%s/%s" % (ilvl.get(qn("w:val")) if ilvl is not None else "?", numId.get(qn("w:val")) if numId is not None else "?"))
    if p.style is not None and p.style.name == "边框代码":
        info.append("STYLE=borders")
    if info:
        out.append("[P%03d] %s :: %r" % (i, " ".join(info), p.text))

out.append("")
out.append("=== SECTIONS PAGE BREAKS ===")
for i, p in enumerate(doc.paragraphs):
    if p._p.pPr is not None:
        for br in p._p.pPr.findall(qn("w:sectPr")):
            out.append("sectPr at P%03d" % i)

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(out))
print("wrote", OUT, len(out))

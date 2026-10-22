# -*- coding: utf-8 -*-
"""Fill the openGauss lab report's gray code boxes with commands from the manual."""
import sys
from docx import Document
from docx.oxml.ns import qn

SRC = sys.argv[1]
DST = sys.argv[2]

doc = Document(SRC)
paras = doc.paragraphs

FILLS = {
    # 1. 切换到 omm 用户
    18: ["su - omm"],
    # 2. 启动数据库服务
    21: ["gs_om -t start"],
    # 3. 连接数据库 postgres，port=26000
    24: ["gsql -d postgres -p 26000 -r"],
    # 4. help 帮助信息
    27: ["help"],
    # 5. 版权信息
    30: ["\\copyright"],
    # 6. openGauss 支持的所有 SQL 语句
    33: ["\\h"],
    # 7. CREATE DATABASE 的详细参数
    36: ["\\help CREATE DATABASE"],
    # 8. 创建数据库 db_tpcc02（SQL 语句以分号结束）
    39: ["CREATE DATABASE db_tpcc02;"],
    # 9. \l 查看数据库列表
    42: ["\\l"],
    # 10. 通过系统表 pg_database 查询数据库列表
    46: ["SELECT datname FROM pg_database;"],
    # 11. 切换数据库至 db_tpcc02（元命令不加分号）
    49: ["\\c db_tpcc02"],
    # 12. 查看创建模式的内容
    52: ["\\h create schema"],
    # 13. 创建自己的模式（张涵怡 -> zhy）
    55: ["CREATE SCHEMA zhy;"],
    # 14. 查看当前数据库中的所有模式
    58: ["\\dn"],
    # 15. 退出数据库连接
    61: ["\\q"],
    # 16. 创建存放命令文件的目录
    65: ["mkdir /home/omm/openGauss"],
    # 17. 用 vi 创建 mysql.sql 并以文件作为命令源
    69: [
        "vi /home/omm/openGauss/mysql.sql",
        "i",
        "select usename , usesysid , usecreatedb from pg_user;",
        "# 按 Esc，输入 :wq 存盘退出",
        "gsql -d postgres -p 26000 -f /home/omm/openGauss/mysql.sql",
    ],
    # 18. 创建数据库用户（张涵怡 -> zhy）
    72: ["CREATE USER zhy IDENTIFIED BY 'Zhy@12345';"],
    # 19. 创建数据库 ST_DB，指定所有者为 zhy
    75: ["CREATE DATABASE st_db OWNER zhy;"],
    # 20. 授予新用户 sysadmin 权限
    78: ["ALTER USER zhy SYSADMIN;"],
    # 21. 用新用户连接 DB_EM 数据库
    81: ["gsql -d db_em -U zhy -p 26000 -W 'Zhy@12345'"],
    # 22. 新建模式 Lab2_scm，设为当前模式并显示
    84: [
        "CREATE SCHEMA lab2_scm;",
        "SET search_path TO lab2_scm;",
        "SHOW search_path;",
    ],
    # 23. 删除数据库 db_temp
    87: ["DROP DATABASE IF EXISTS db_temp;"],
    # 24. 职级表 rank 建表语句
    121: [
        "CREATE TABLE rank",
        "(",
        "    rankid     INTEGER     NOT NULL,",
        "    rankname   VARCHAR(50) NOT NULL,",
        "    rankremark TEXT",
        ");",
    ],
    # 25. 部门表 department 建表语句
    123: [
        "CREATE TABLE department",
        "(",
        "    departmentid     INTEGER     NOT NULL,",
        "    departmentname   VARCHAR(50) NOT NULL,",
        "    departmentremark TEXT",
        ");",
    ],
    # 26. 职员表 people 建表语句
    126: [
        "CREATE TABLE people",
        "(",
        "    peopleid      INTEGER                       NOT NULL,",
        "    departmentid  INTEGER                       NOT NULL,",
        "    rankid        INTEGER                       NOT NULL,",
        "    peoplename    VARCHAR(50)                   NOT NULL,",
        "    peoplesex     VARCHAR(4)                    NOT NULL,",
        "    peoplebirth   TIMESTAMP(0) WITHOUT TIME ZONE NOT NULL,",
        "    peoplesalary  NUMERIC(12,2)                 NOT NULL,",
        "    peoplephone   VARCHAR(20)                   NOT NULL,",
        "    peopleaddress VARCHAR(100),",
        "    peopleaddtime SMALLDATETIME",
        ");",
    ],
}

PEOPLE = [
    (1, 1, 1, "赵小明", "男", "1970-01-01", 50000.00, "13893330005", "北京"),
    (2, 1, 1, "司马相如", "男", "1968-02-09", 40000.00, "13554785452", "成都"),
    (3, 1, 1, "周小顺", "女", "1975-03-09", 25000.00, "13256854578", "甘肃"),
    (4, 1, 1, "钱小旦", "女", "1973-04-08", 12000.00, "13985745871", "山东"),
    (5, 2, 1, "张世同", "男", "1980-05-08", 18000.00, "13535987412", "安徽"),
    (6, 2, 3, "委正大", "男", "1978-06-08", 9000.00, "13845789568", "河北"),
    (7, 2, 2, "赵有飞", "男", "1985-07-08", 10500.00, "13878562568", "浙江"),
    (8, 2, 2, "李小强", "男", "1983-08-20", 8500.00, "13335457412", "山西"),
    (9, 3, 1, "胡瑗瑗", "女", "1990-09-20", 16500.00, "13437100050", "北京"),
    (10, 3, 2, "宋可睛", "女", "1988-12-20", 15000.00, "13889562354", "北京"),
    (11, 3, 3, "肖飞", "男", "1995-11-11", 9000.00, "13385299632", "河南"),
    (12, 3, 1, "张锐杰", "男", "1993-10-13", 18000.00, "13478545263", "上海"),
    (13, 4, 1, "曹德旺", "男", "2000-12-12", 17500.00, "13878523695", "深圳"),
    (14, 4, 3, "李萌", "女", "1998-11-22", 9000.00, "13698545841", "广州"),
    (15, 4, 3, "叶子飞云", "女", "1994-01-22", 11000.00, "13558745874", "内蒙古"),
    (16, 4, 2, "孙大勇", "男", "1993-02-21", 12000.00, "13698745214", "南京"),
    (17, 5, 2, "吴若欣", "女", "1992-02-21", 13000.00, "13985478512", "湖南"),
    (18, 5, 1, "刘婷", "女", "1995-02-21", 13500.00, "13778787874", "湖北"),
    (19, 5, 2, "周海涛", "男", "1996-10-11", 8000.00, "13987455214", "天津"),
    (20, 5, 3, "韦小宝", "男", "1997-09-10", 5500.00, "13254785965", "宁夏"),
    (21, 1, 3, "王萌", "女", "1986-05-19", 8500.00, "13352197364", "新疆"),
    (22, 2, 1, "陆子虹", "女", "1998-05-19", 7500.00, "13025457392", "南京"),
    (23, 3, 2, "康军", "男", "1999-06-01", 7500.00, "13077778888", "广西"),
]

insert_lines = [
    "-- 职级表 rank",
    "INSERT INTO rank",
    "    (rankid, rankname, rankremark)",
    "VALUES",
    "(1, '初级', '获得HCIA证书，具备行业项目的开发的基本水平'),",
    "(2, '中级', '获得HCIP证书,具备横向项目独立设计的能力'),",
    "(3, '高级', '获得HCIE相当的证书具备综合项目设计丰富经历');",
    "",
    "-- 部门表 department",
    "INSERT INTO department",
    "    (departmentid, departmentname, departmentremark)",
    "VALUES",
    "(1, '软件部', '......'),",
    "(2, '硬件部', '......'),",
    "(3, '市场部', '......'),",
    "(4, '财务部', '企业财务结算'),",
    "(5, '人力部', '人事管理'),",
    "(6, '品控部', '质量检控');",
    "",
    "-- 职员表 people",
    "INSERT INTO people",
    "    (peopleid, departmentid, rankid, peoplename, peoplesex,",
    "     peoplebirth, peoplesalary, peoplephone, peopleaddress)",
    "VALUES",
]
for idx, row in enumerate(PEOPLE):
    line = "(%d,%d,%d,'%s','%s','%s',%.2f,'%s','%s')" % row
    line += ";" if idx == len(PEOPLE) - 1 else ","
    insert_lines.append(line)
FILLS[178] = insert_lines


def style_run(run):
    """Use a widely available monospace font; the template's Huawei Sans is not installed."""
    run.font.name = "Consolas"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "宋体")
    rFonts.set(qn("w:cs"), "Consolas")


def set_lines(paragraph, lines):
    # remove existing runs but keep paragraph properties
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(lines[0])
    style_run(run)
    for line in lines[1:]:
        run.add_break()
        run.add_text(line)


for idx, lines in FILLS.items():
    set_lines(paras[idx], lines)

# header fields
if paras[1].runs[1].text.strip("_ ") == "":
    paras[1].runs[1].text = "张涵怡"
if paras[2].runs[1].text.strip("_ ") == "":
    paras[2].runs[1].text = "2026-10-20"

doc.save(DST)
print("saved", DST)
